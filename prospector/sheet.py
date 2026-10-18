"""Prospects.xlsx: companies, contacts, emails and Apollo credit use."""
import json

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from . import config, db


def _sheet(wb, title: str, headers: list[str], rows: list[list]) -> None:
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
    for r in rows:
        ws.append(r)
    for i, h in enumerate(headers, 1):
        width = max([len(h)] + [len(str(r[i - 1] or "")) for r in rows])
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 55)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def export() -> str:
    wb = Workbook()
    wb.remove(wb.active)
    contacts = db.q("""SELECT ct.*, co.name company, co.domain FROM contacts ct JOIN companies co ON co.id = ct.company_id
                       ORDER BY COALESCE(ct.sent_at, ct.created_at) DESC""")
    _sheet(wb, "Contacts", ["Name", "Title", "Company", "Email", "Who", "Status", "Sent", "Follow-up", "Replied",
                            "Reply"],
           [[r["name"], r["title"], r["company"], r["email"], config.TIERS.get(r["tier"], ("",))[0], r["status"],
             (r["sent_at"] or "")[:16].replace("T", " "), (r["followup_at"] or "")[:16].replace("T", " "),
             (r["replied_at"] or "")[:16].replace("T", " "), r["reply_snippet"]] for r in contacts])
    companies = db.q("SELECT * FROM companies ORDER BY id DESC")
    _sheet(wb, "Companies", ["Company", "Website", "Size", "Open roles", "Source", "Status", "Added"],
           [[r["name"], r["website"] or r["domain"], {1: "1-50", 2: "51-200", 3: "201-500"}.get(r["size_stage"], ""),
             ", ".join(x["title"] for x in json.loads(r["roles_json"] or "[]")), r["source"], r["status"],
             (r["created_at"] or "")[:10]] for r in companies])
    msgs = db.q("""SELECT m.*, ct.name, ct.email, co.name company FROM messages m JOIN contacts ct ON ct.id = m.contact_id
                   JOIN companies co ON co.id = ct.company_id WHERE m.status='sent' ORDER BY m.sent_at DESC""")
    _sheet(wb, "Emails sent", ["Date", "To", "Email", "Company", "Type", "Subject", "Body"],
           [[(r["sent_at"] or "")[:16].replace("T", " "), r["name"], r["email"], r["company"], r["kind"], r["subject"],
             r["body"]] for r in msgs])
    usage = db.q("""SELECT substr(created_at,1,7) month, endpoint, SUM(credits) credits FROM apollo_usage
                    GROUP BY month, endpoint ORDER BY month DESC""")
    _sheet(wb, "Apollo credits", ["Month", "Used for", "Credits"], [[r["month"], r["endpoint"], r["credits"]] for r in usage])
    try:
        wb.save(config.SHEET_PATH)
        return str(config.SHEET_PATH)
    except PermissionError:  # open in Excel
        alt = config.SHEET_PATH.with_name(f"Prospects_{db.now().replace(':', '')}.xlsx")
        wb.save(alt)
        return str(alt)


def stats_text() -> str:
    c = {r["status"]: r["n"] for r in db.q("SELECT status, COUNT(*) n FROM contacts GROUP BY status")}
    companies = db.q1("SELECT COUNT(*) n FROM companies")["n"]
    sent = db.q1("SELECT COUNT(*) n FROM messages WHERE status='sent' AND kind='initial'")["n"]
    fups = db.q1("SELECT COUNT(*) n FROM messages WHERE status='sent' AND kind='followup'")["n"]
    replied = c.get("replied", 0)
    rate = f" ({100 * replied // sent}%)" if sent else ""
    return (f"🏢 Companies: {companies}   👤 People found: {sum(c.values())}\n"
            f"📨 Emails sent: {sent}  ·  follow-ups: {fups}\n"
            f"💬 Replies: {replied}{rate}   ↩️ Bounced: {c.get('bounced', 0)}\n"
            f"⏳ Ready to email: {c.get('revealed', 0)}   🔎 Waiting for email: {c.get('found', 0)}\n"
            f"🪙 Apollo credits this month: {db.credits_this_month()} / {config.APOLLO_MONTHLY_CREDITS}\n"
            f"📏 Company size: stage {db.get_setting('stage', '1')}   ⏸ Paused: {'yes' if db.get_setting('paused') == '1' else 'no'}")
