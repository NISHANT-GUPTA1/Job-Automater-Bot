"""Stats and Excel export of everything tracked."""
import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from . import config, db


def stats_text() -> str:
    rows = db.q("SELECT status, COUNT(*) c FROM jobs GROUP BY status ORDER BY c DESC")
    total = sum(r["c"] for r in rows)
    today = datetime.now().date().isoformat()
    applied_today = db.q1("SELECT COUNT(*) c FROM jobs WHERE applied_at LIKE ?", (today + "%",))["c"]
    emails = db.q1("SELECT COUNT(*) c FROM emails WHERE status='sent'")["c"]
    emails_today = db.q1("SELECT COUNT(*) c FROM emails WHERE status='sent' AND sent_at LIKE ?", (today + "%",))["c"]
    accounts = db.q1("SELECT COUNT(*) c FROM accounts")["c"]
    learned = db.q1("SELECT COUNT(*) c FROM qa_memory")["c"]
    lines = [f"📊 Tracked jobs: {total}"] + [f"   • {r['status']}: {r['c']}" for r in rows]
    lines += [f"✅ Applied today: {applied_today}",
              f"📨 Cold emails sent: {emails} (today {emails_today})",
              f"🔐 Site accounts: {accounts}   🧠 Learned answers: {learned}",
              "", cost_text()]
    return "\n".join(lines)


def cost_text() -> str:
    rows = db.q("""SELECT provider, COUNT(*) calls, SUM(input_tokens) i, SUM(output_tokens) o,
                   SUM(cache_read) cr, SUM(cost_usd) cost FROM llm_usage GROUP BY provider""")
    if not rows:
        return "💸 No LLM calls yet."
    lines = ["💸 LLM usage:"]
    for r in rows:
        lines.append(f"   • {r['provider']}: {r['calls']} calls, in {r['i'] or 0:,} / out {r['o'] or 0:,} tok"
                     f", cached {r['cr'] or 0:,}" + (f", ${r['cost']:.4f}" if r["cost"] else ""))
    per_job = db.q1("""SELECT AVG(s) a FROM (SELECT SUM(cost_usd) s FROM llm_usage
                       WHERE provider='anthropic' AND job_id IS NOT NULL GROUP BY job_id)""")["a"]
    if per_job:
        lines.append(f"   • Avg Anthropic cost per job: ${per_job:.4f}")
    return "\n".join(lines)


def job_cost(job_id: int) -> dict:
    r = db.q1("""SELECT COUNT(*) calls, COALESCE(SUM(input_tokens),0) i, COALESCE(SUM(output_tokens),0) o,
                 COALESCE(SUM(cache_read),0) cr, COALESCE(SUM(cost_usd),0) cost
                 FROM llm_usage WHERE job_id=? AND provider='anthropic'""", (job_id,))
    other = db.q1("SELECT COUNT(*) c FROM llm_usage WHERE job_id=? AND provider!='anthropic'", (job_id,))["c"]
    return {"calls": r["calls"], "in": r["i"], "out": r["o"], "cached": r["cr"], "usd": r["cost"], "fallback_calls": other}


def job_cost_text(job_id: int) -> str:
    c = job_cost(job_id)
    inr = c["usd"] * float(config.env("USD_TO_INR", "88"))
    text = (f"💸 Claude cost for job #{job_id}: ${c['usd']:.4f} (≈ ₹{inr:.2f}) — {c['calls']} calls, "
            f"{c['in'] + c['cached']:,} in ({c['cached']:,} cached) / {c['out']:,} out tokens")
    if c["fallback_calls"]:
        text += f"\n   + {c['fallback_calls']} fallback calls (OpenAI/Gemini/Groq)"
    return text


def _sheet(wb, title, headers, rows):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
    for r in rows:
        ws.append([r[h] if h in r.keys() else "" for h in headers])
    for i, h in enumerate(headers, 1):
        width = max([len(h)] + [len(str(r[h] or "")) for r in rows])
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 62)
    ws.freeze_panes = "A2"
    return ws


MAILED_PATH = config.DATA_DIR / "Mailed_Contacts.xlsx"
MAILED_HEADERS = ["#", "Date sent", "HR / Contact name", "Email", "Company", "Position", "Subject",
                  "Source", "Job #", "Resume sent"]


def mailed_rows() -> list[dict]:
    """Every cold email actually sent, newest first, with who/where/what."""
    rows = []
    for r in db.q("""SELECT e.*, j.title, j.company, j.url, j.details_json FROM emails e
                     LEFT JOIN jobs j ON j.id = e.job_id WHERE e.status='sent' ORDER BY e.sent_at DESC"""):
        d = json.loads(r["details_json"] or "{}")
        source = ("Outreach" if d.get("outreach") else
                  "Screenshot / pasted post" if (r["url"] or "").startswith("text:") else "Job link")
        rows.append({"#": 0, "Date sent": (r["sent_at"] or "")[:16].replace("T", " "),
                     "HR / Contact name": d.get("person_name") or d.get("contact_name") or "",
                     "Email": r["to_addr"], "Company": r["company"] or "", "Position": r["title"] or "",
                     "Subject": r["subject"], "Source": source, "Job #": r["job_id"],
                     "Resume sent": Path(r["attachment"] or "").name})
    for i, row in enumerate(rows, 1):
        row["#"] = len(rows) - i + 1
    return rows


def update_mailed_sheet() -> str:
    """(Re)write data/Mailed_Contacts.xlsx - called after every sent email, and by /mailed."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Mailed contacts"
    ws.append(MAILED_HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
    rows = mailed_rows()
    for r in rows:
        ws.append([r[h] for h in MAILED_HEADERS])
    for i, h in enumerate(MAILED_HEADERS, 1):
        width = max([len(h)] + [len(str(r[h] or "")) for r in rows])
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 50)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    try:
        wb.save(MAILED_PATH)
    except PermissionError:  # file is open in Excel - save a copy instead of failing
        alt = MAILED_PATH.with_name(f"Mailed_Contacts_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
        wb.save(alt)
        return str(alt)
    return str(MAILED_PATH)


ACCOUNTS_PATH = config.DATA_DIR / "Site_Accounts.xlsx"
ACCOUNT_HEADERS = ["Company", "Website", "Login URL", "Email", "Password", "Created", "Job #"]


def account_rows() -> list[dict]:
    rows = []
    for a in db.list_accounts():
        r = db.q1("SELECT j.company, j.id FROM accounts ac LEFT JOIN jobs j ON j.id = ac.job_id WHERE ac.domain=?",
                  (a["domain"],))
        rows.append({"Company": (r["company"] if r else "") or "", "Website": a["domain"], "Login URL": a["login_url"],
                     "Email": a["email"], "Password": a["password"], "Created": (a["created_at"] or "")[:16].replace("T", " "),
                     "Job #": r["id"] if r else ""})
    return rows


def update_accounts_sheet() -> str:
    """(Re)write data/Site_Accounts.xlsx: every job-site account the bot created, with its password."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Site accounts"
    ws.append(ACCOUNT_HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
    rows = account_rows()
    for r in rows:
        ws.append([r[h] for h in ACCOUNT_HEADERS])
    for i, h in enumerate(ACCOUNT_HEADERS, 1):
        width = max([len(h)] + [len(str(r[h] or "")) for r in rows])
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 60)
    ws.freeze_panes = "A2"
    try:
        wb.save(ACCOUNTS_PATH)
    except PermissionError:  # open in Excel
        alt = ACCOUNTS_PATH.with_name(f"Site_Accounts_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
        wb.save(alt)
        return str(alt)
    return str(ACCOUNTS_PATH)


def export_excel() -> str:
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "Jobs", ["id", "status", "title", "company", "location", "match_score", "url", "contact_email",
                        "applied_at", "emailed_at", "resume_path", "claude_cost_usd", "notes", "summary", "created_at"],
           db.q("""SELECT j.*, ROUND((SELECT SUM(cost_usd) FROM llm_usage u WHERE u.job_id=j.id), 5) claude_cost_usd
                   FROM jobs j ORDER BY id DESC"""))
    _sheet(wb, "Mailed contacts", MAILED_HEADERS, mailed_rows())
    _sheet(wb, "Emails", ["id", "job_id", "status", "to_addr", "subject", "body", "attachment", "sent_at", "error"],
           db.q("SELECT * FROM emails WHERE status IN ('sent','failed') ORDER BY id DESC"))
    _sheet(wb, "Accounts", ["domain", "email", "login_url", "created_at"],
           db.q("SELECT domain, email, login_url, created_at FROM accounts ORDER BY id DESC"))
    _sheet(wb, "Learned answers", ["id", "question", "answer", "times_used"],
           db.q("SELECT * FROM qa_memory ORDER BY id"))
    _sheet(wb, "Job descriptions", ["id", "title", "company", "description"],
           db.q("SELECT id, title, company, description FROM jobs ORDER BY id DESC"))
    path = config.DATA_DIR / f"job_tracker_{datetime.now():%Y%m%d_%H%M}.xlsx"
    wb.save(path)
    return str(path)
